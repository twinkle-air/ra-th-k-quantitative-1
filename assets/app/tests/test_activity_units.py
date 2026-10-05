"""Independent dimensional regression: do not lock the faulty legacy divisor."""
import math
import unittest
from io import BytesIO

from openpyxl import load_workbook
from pypdf import PdfReader
from app.services.activity_units import CONTENT_DIVISORS, activity_to_content, bq_kg_per_ppm
from app.services.exporters import _display_content, _reported_content, _residual_rows, export_xlsx, export_pdf, export_png
from app.services.report_templates import _sample_values, LANGUAGE_LABELS


class ActivityUnitTests(unittest.TestCase):
    def test_independent_one_milligram_definition_for_ra_and_th(self):
        # A = lambda * N, with one milligram of isotope per kg, independent of production helper.
        for isotope, years, mass in [('Ra226', 1600, 226.0254), ('Th232', 1.405e10, 232.03806)]:
            atoms = 0.001 / mass * 6.02214076e23
            expected = atoms * math.log(2) / (years * 31557600)
            self.assertAlmostEqual(bq_kg_per_ppm(years, mass) / expected, 1, places=12)
            tolerance = 1e-12 if isotope == 'Ra226' else .001
            self.assertLess(abs(CONTENT_DIVISORS[isotope] / expected - 1), tolerance)
        self.assertAlmostEqual(activity_to_content(73.01628, 'Ra226'), 1.996294157556173e-6, places=16)
        self.assertAlmostEqual(activity_to_content(73.01628, 'Ra226') * 1000, .001996294157556173, places=15)

    def test_corrected_conversion_all_presentations_and_formal_template_gate(self):
        from test_core import synthetic, AnalysisSettings
        from app.services.analysis import analyze_batch
        result = analyze_batch(synthetic('standard', 1), [synthetic('sample', .3)], [244], AnalysisSettings(),
                               {'calibration': {'slope': .2, 'intercept': .1}, 'sample': {'slope': .2, 'intercept': .1}})
        row = result['results'][0]
        # Manual calibration intentionally has no fitted points. Supply generator-known
        # synthetic points solely to exercise residual presentation, not scientific validation.
        row['calibration']['matched_points'] = [[100, 20.1], [200, 40.1]]
        expected = row['activity_bq_kg']['Ra226'] / CONTENT_DIVISORS['Ra226']
        self.assertAlmostEqual(row['ra_ppm'], expected)
        self.assertAlmostEqual(_display_content(row, 'Ra226'), expected)
        self.assertIsNone(_reported_content(row, 'Ra226'))
        self.assertEqual(_sample_values(row, LANGUAGE_LABELS['en'])['sample.ra_ppm'], '—')
        for language in ('zh', 'zht', 'en', 'fr'):
            book = load_workbook(BytesIO(export_xlsx(result, language)), data_only=True)
            self.assertAlmostEqual(book.worksheets[0]['C4'].value, expected)
            self.assertIn('conditional_unmatched_geometry', book.worksheets[0]['A6'].value)
            self.assertIn('u_counting', book.sheetnames)
            self.assertGreater(book['Calibration Residuals'].max_row, 1)
            self.assertTrue(export_png(result, language).startswith(b'\x89PNG'))
            text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(export_pdf(result, language))).pages)
            self.assertIn('conditional_unmatched_geometry', text)
            self.assertIn(f'{expected:.3g}', text)
        row['quality']['nuclides']['Ra226']['reportable_activity_bq_kg'] = row['activity_bq_kg']['Ra226']
        self.assertAlmostEqual(float(_sample_values(row, LANGUAGE_LABELS['en'])['sample.ra_ppm']), expected, delta=expected*.001)

    def test_signed_residuals_not_correlation_proxy(self):
        row = {'calibration': {'slope': .2, 'intercept': -.1, 'matched_points': [[100, 20], [200, 39.8]]}}
        residuals = _residual_rows(row)
        self.assertAlmostEqual(residuals[0][2], -.1)
        self.assertAlmostEqual(residuals[1][2], .1)
        self.assertAlmostEqual(residuals[0][3], .1)
