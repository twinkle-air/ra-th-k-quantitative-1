from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
import re
from typing import Any


SUPPORTED_PARAMETER_EXTENSIONS = {".pdf", ".xls", ".xlsx"}
MAX_PARAMETER_FILE_SIZE = 20 * 1024 * 1024


def _excel_text(filename: str, payload: bytes) -> str:
    extension = Path(filename).suffix.lower()
    rows: list[list[Any]] = []
    if extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            rows.append([f"[{sheet.title}]"])
            rows.extend([list(row) for row in sheet.iter_rows(values_only=True)])
    else:
        import xlrd

        workbook = xlrd.open_workbook(file_contents=payload)
        for sheet in workbook.sheets():
            rows.append([f"[{sheet.name}]"])
            rows.extend(sheet.row_values(index) for index in range(sheet.nrows))
    return "\n".join(" : ".join(str(cell).strip() for cell in row if cell not in (None, "")) for row in rows)


def _pdf_text(payload: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("读取 PDF 参数文件需要安装 pypdf。") from exc
    reader = PdfReader(BytesIO(payload))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - library-specific failures
            raise ValueError("PDF 已加密，无法读取参数。") from exc
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if not text.strip():
        raise ValueError("PDF 中没有可提取文字；扫描版 PDF 请先进行 OCR。")
    return text


def _normalize(text: str) -> str:
    translations = str.maketrans({"²": "2", "³": "3", "⁴": "4", "⁰": "0", "⁶": "6", "：": ":", "，": ",", "％": "%", "−": "-"})
    return re.sub(r"[\t\u00a0 ]+", " ", text.translate(translations)).replace("\r", "")


NUMBER = r"([-+]?\d+(?:[.,]\d+)?)"
FIELD_PATTERNS: dict[str, tuple[str, ...]] = {
    "calibration_mass_g": (
        r"(?:校准源|校準源|标准源|標準源)(?:净|淨)?(?:质量|質量|重量)",
        r"(?:calibration|standard|reference)\s+(?:source\s+)?mass",
        r"masse\s+(?:de\s+)?(?:la\s+)?source",
    ),
    "ra_activity_bq": (r"(?:226\s*Ra|Ra\s*[-–]?\s*226|镭\s*[-–]?\s*226|鐳\s*[-–]?\s*226)(?:\s*(?:活度|activity|activité))?",),
    "th_activity_bq": (r"(?:232\s*Th|Th\s*[-–]?\s*232|钍\s*[-–]?\s*232|釷\s*[-–]?\s*232)(?:\s*(?:活度|activity|activité))?",),
    "k_activity_bq": (r"(?:40\s*K|K\s*[-–]?\s*40|钾\s*[-–]?\s*40|鉀\s*[-–]?\s*40)(?:\s*(?:活度|activity|activité))?",),
    "roi_half_width_keV": (r"ROI\s*(?:半宽|半寬|half\s*[- ]?width|demi\s*[- ]?largeur)",),
    "background_gap_keV": (r"(?:本底|背景|background|fond)\s*(?:间隔|間隔|gap|écart)",),
    "background_width_keV": (r"(?:本底|背景|background|fond)\s*(?:窗宽|窗寬|window\s*width|largeur\s*(?:de\s+la\s+)?fenêtre)",),
    "calibration_slope": (r"(?:刻度|校准|校準|calibration|étalonnage)?\s*(?:斜率|slope|pente)\b",),
    "calibration_intercept": (r"(?:刻度|校准|校準|calibration|étalonnage)?\s*(?:截距|intercept|ordonnée\s+à\s+l['’]origine)\b",),
}


def _number_after(text: str, aliases: tuple[str, ...]) -> tuple[float, str] | None:
    for alias in aliases:
        match = re.search(rf"({alias})[^\d+\-]{{0,35}}{NUMBER}\s*(kg|g|Bq|keV|%)?", text, re.IGNORECASE)
        if not match:
            continue
        value = float(match.group(2).replace(",", "."))
        unit = (match.group(3) or "").lower()
        label = re.sub(r"\s+", " ", match.group(1)).strip()
        return value, f"{label}{' (' + unit + ')' if unit else ''}"
    return None


def _reference_date(text: str) -> tuple[str, str] | None:
    labels = r"(?:活度)?参考日期|(?:活度)?參考日期|reference\s+date|activity\s+reference\s+date|date\s+de\s+référence"
    match = re.search(rf"({labels})[^\d]{{0,30}}(\d{{4}}[年/\-.]\d{{1,2}}[月/\-.]\d{{1,2}}日?|\d{{1,2}}[/\-.]\d{{1,2}}[/\-.]\d{{4}})", text, re.IGNORECASE)
    if not match:
        return None
    raw = match.group(2).replace("年", "-").replace("月", "-").replace("日", "").replace("/", "-").replace(".", "-")
    formats = ("%Y-%m-%d", "%d-%m-%Y")
    for fmt in formats:
        try:
            return datetime.strptime(raw, fmt).strftime("%Y-%m-%d"), match.group(1)
        except ValueError:
            continue
    return None


def _boolean_after(text: str, aliases: str) -> tuple[bool, str] | None:
    match = re.search(rf"({aliases})[^\n]{{0,45}}", text, re.IGNORECASE)
    if not match:
        return None
    phrase = match.group(0).lower()
    if re.search(r"未确认|未確認|否|false|no|non|unchecked|not\s+confirmed", phrase):
        return False, match.group(1)
    if re.search(r"已确认|已確認|是|true|yes|oui|checked|confirmed|appliqu", phrase):
        return True, match.group(1)
    return None


def parse_analysis_parameters(filename: str, payload: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_PARAMETER_EXTENSIONS:
        raise ValueError("参数文件仅支持 .pdf、.xls 和 .xlsx。")
    if not payload:
        raise ValueError("参数文件为空。")
    if len(payload) > MAX_PARAMETER_FILE_SIZE:
        raise ValueError("参数文件不能超过 20 MB。")
    raw_text = _pdf_text(payload) if extension == ".pdf" else _excel_text(filename, payload)
    text = _normalize(raw_text)
    values: dict[str, Any] = {}
    sources: dict[str, str] = {}
    for key, aliases in FIELD_PATTERNS.items():
        found = _number_after(text, aliases)
        if not found:
            continue
        value, source = found
        if key == "calibration_mass_g" and "kg" in source.lower():
            value *= 1000
        values[key] = value
        sources[key] = source
    date = _reference_date(text)
    if date:
        values["reference_date"], sources["reference_date"] = date
    equilibrium = _boolean_after(text, r"(?:Ra\s*/\s*Th\s*)?(?:衰变链|衰變鏈|chain)\s*(?:平衡|equilibrium|équilibre)")
    if equilibrium:
        values["assume_chain_equilibrium"], sources["assume_chain_equilibrium"] = equilibrium
    correction = _boolean_after(text, r"(?:K\s*[-–]?\s*40|40\s*K).{0,20}(?:干扰|干擾|interference|interférence).{0,20}(?:修正|correction|corriger)")
    if correction:
        values["correct_k_interference"], sources["correct_k_interference"] = correction
    if not values:
        raise ValueError("未识别到可用分析参数；请检查字段名称或改用可提取文字的文件。")
    return {"filename": filename, "values": values, "sources": sources, "recognized_count": len(values)}
