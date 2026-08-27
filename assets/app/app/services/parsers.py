from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO, StringIO
from pathlib import Path
import csv
import re
from typing import Iterable

import numpy as np

from .word_documents import read_word_rows


SUPPORTED_EXTENSIONS = {".xls", ".xlsx", ".txt", ".csv", ".dat", ".doc", ".docx"}


@dataclass
class Spectrum:
    name: str
    channels: np.ndarray
    counts: np.ndarray
    live_time: float
    real_time: float | None = None
    acquired_at: datetime | None = None
    metadata: dict[str, str] = field(default_factory=dict)

    @property
    def dead_time_fraction(self) -> float | None:
        if not self.real_time or self.real_time <= 0:
            return None
        return max(0.0, 1.0 - self.live_time / self.real_time)


def _to_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return None


def _read_xlsx(payload: bytes) -> list[list[object]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _read_xls(payload: bytes) -> list[list[object]]:
    try:
        import xlrd
    except ImportError as exc:
        raise ValueError("读取 .xls 需要安装 xlrd（pip install xlrd）。") from exc
    book = xlrd.open_workbook(file_contents=payload)
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(index) for index in range(sheet.nrows)]


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文本编码；请转换为 UTF-8 或 GB18030。")


def _read_text(payload: bytes) -> list[list[object]]:
    text = _decode_text(payload)
    lines = [line for line in text.splitlines() if line.strip()]
    if not lines:
        return []
    sample = "\n".join(lines[:20])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t; ")
        return [row for row in csv.reader(StringIO(text), dialect)]
    except csv.Error:
        return [re.split(r"[\s,;\t]+", line.strip()) for line in lines]


def _read_rows(filename: str, payload: bytes) -> list[list[object]]:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持 {extension or '无扩展名'} 文件。")
    if len(payload) > 50 * 1024 * 1024:
        raise ValueError("单个文件不能超过 50 MB。")
    if extension == ".xlsx":
        return _read_xlsx(payload)
    if extension == ".xls":
        return _read_xls(payload)
    if extension in {".doc", ".docx"}:
        return read_word_rows(filename, payload)
    return _read_text(payload)


def _canonical_metadata_key(key: str) -> str:
    compact = re.sub(r"[\s_\-（）()\[\]：:/.]+", "", key).upper()
    if compact.startswith(("TLIVE", "LIVETIME")) or compact == "有效活时间" or "活时间" in compact:
        return "TLIVE"
    if compact.startswith(("TREAL", "REALTIME")) or compact in {"实时间", "实际时间", "总测量时间"}:
        return "TREAL"
    if compact.startswith("DEADTIME") or compact in {"死时间", "死时间率", "死时间百分比"}:
        return "DEADTIME"
    return re.sub(r"\s+", "_", key.strip()).upper()


def _extract_metadata(rows: Iterable[list[object]]) -> dict[str, str]:
    metadata: dict[str, str] = {}
    for row in rows:
        if not row:
            continue
        first = str(row[0]).strip()
        if "=" in first:
            key, value = first.split("=", 1)
            metadata[_canonical_metadata_key(key)] = value.strip()
        elif len(row) > 1 and _to_float(row[0]) is None:
            key = _canonical_metadata_key(first)
            metadata[key] = str(row[1]).strip()
    return metadata


def _live_time_from_metadata(metadata: dict[str, str]) -> tuple[float | None, str | None]:
    direct = _to_float(metadata.get("TLIVE"))
    if direct is not None and direct > 0:
        return direct, "TLIVE"
    real_time = _to_float(metadata.get("TREAL"))
    dead_time = _to_float(metadata.get("DEADTIME"))
    if real_time is not None and real_time > 0 and dead_time is not None and dead_time >= 0:
        fraction = dead_time / 100 if dead_time > 1 else dead_time
        if 0 <= fraction < 1:
            return real_time * (1 - fraction), "TREAL+DEADTIME"
    return None, None


def inspect_spectrum_metadata(filename: str, payload: bytes) -> dict[str, object]:
    rows = _read_rows(filename, payload)
    metadata = _extract_metadata(rows)
    live_time, source = _live_time_from_metadata(metadata)
    real_time = _to_float(metadata.get("TREAL"))
    return {
        "live_time_s": live_time,
        "live_time_source": source,
        "real_time_s": real_time,
        "acquired_at": _parse_datetime(metadata).isoformat() if _parse_datetime(metadata) else None,
        "metadata": metadata,
    }


def _parse_datetime(metadata: dict[str, str]) -> datetime | None:
    date = metadata.get("DATE")
    time = metadata.get("TIME", "00:00:00")
    if not date:
        return None
    value = f"{date} {time}"
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def parse_spectrum(filename: str, payload: bytes, live_time_override: float | None = None) -> Spectrum:
    rows = _read_rows(filename, payload)

    metadata = _extract_metadata(rows)
    numeric: list[tuple[float, float]] = []
    for row in rows:
        if len(row) < 2:
            continue
        channel, count = _to_float(row[0]), _to_float(row[1])
        if channel is not None and count is not None and channel >= 0 and count >= 0:
            numeric.append((channel, count))
    if len(numeric) < 32:
        raise ValueError("没有找到足够的“道址—计数”数据（至少需要 32 行）。")

    numeric.sort(key=lambda item: item[0])
    channels = np.asarray([item[0] for item in numeric], dtype=float)
    counts = np.asarray([item[1] for item in numeric], dtype=float)
    if np.any(np.diff(channels) <= 0):
        unique, indices = np.unique(channels, return_index=True)
        channels, counts = unique, counts[indices]

    detected_live_time, _ = _live_time_from_metadata(metadata)
    live_time = live_time_override or detected_live_time
    if not live_time or live_time <= 0:
        raise ValueError("文件未提供有效 TLIVE；请在导入时填写活时间（秒）。")
    real_time = _to_float(metadata.get("TREAL"))
    return Spectrum(
        name=Path(filename).stem,
        channels=channels,
        counts=counts,
        live_time=float(live_time),
        real_time=real_time,
        acquired_at=_parse_datetime(metadata),
        metadata=metadata,
    )
